"""Import de produits par fichier Excel (.xlsx) ou CSV : outil des forfaits Pro et Business.

Deux temps : un aperçu (rien n'est créé) puis la création des lignes valides. Chaque ligne passe par les mêmes
règles que le formulaire « Ajouter un produit » (ProductCreateSerializer), la limite de produits du forfait et la
limite alimentaire. Un produit déjà présent sous le même nom n'est pas recréé. Les photos s'ajoutent ensuite.
"""
import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from rest_framework import permissions, status
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.response import Response

from payments.plan_features import PlanFeatureRequired, has, require
from payments.subscription_check import SubscriptionChecker
from products.models import Product, ProductCategory
from products.serializers import ProductCreateSerializer
from stores.models import Store

MAX_ROWS = 500
MAX_BYTES = 2 * 1024 * 1024

# Colonne -> noms acceptés dans l'en-tête (sans accents, sans majuscules).
COLUMNS = {
    'name': ('nom', 'nom du produit', 'produit', 'designation', 'name'),
    'price': ('prix', 'prix de vente', 'prix fcfa', 'prix f', 'price'),
    'stock': ('stock', 'quantite', 'qte', 'quantite en stock'),
    'weight_kg': ('poids', 'poids kg', 'poids en kg', 'weight'),
    'length_m': ('longueur', 'longueur m', 'longueur en m', 'length'),
    'category': ('categorie', 'category', 'rayon'),
    'description': ('description', 'details', 'detail'),
    'compare_price': ('prix avant promo', 'ancien prix', 'prix barre', 'prix initial'),
    'brand': ('marque', 'brand'),
    'model': ('modele', 'model'),
    'color': ('couleur', 'couleurs', 'color'),
    'sizes': ('tailles', 'taille', 'pointures', 'sizes'),
    'sku': ('reference', 'ref', 'sku'),
    'barcode': ('code barres', 'code barre', 'ean', 'barcode'),
}
REQUIRED = ('name', 'price', 'weight_kg', 'length_m')
TEMPLATE_HEADER = ['Nom', 'Prix', 'Stock', 'Poids (kg)', 'Longueur (m)', 'Catégorie', 'Description',
                   'Prix avant promo', 'Marque', 'Couleur', 'Tailles', 'Référence', 'Code-barres']
TEMPLATE_ROWS = [
    ['Riz parfumé 5 kg', '4500', '20', '5', '0,4', '', 'Sac de 5 kg', '', 'Uncle Bens', '', '', 'RIZ5', ''],
    ['Baskets en toile', '12000', '6', '0,8', '0,35', '', 'Confortables', '15000', 'Converse', 'Noir', '39, 40, 41', '', ''],
]


def _plain(value):
    text = unicodedata.normalize('NFD', str(value or '').strip().lower())
    text = re.sub(r'[̀-ͯ]', '', text)
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', text).split())


ALIASES = {_plain(alias): key for key, names in COLUMNS.items() for alias in names}


def _error(code, message, next_step, http=status.HTTP_400_BAD_REQUEST):
    return Response({'success': False, 'error': {'code': code, 'message': message, 'reason': message,
                                                 'next_step': next_step}}, status=http)


def _store(request):
    return Store.objects.select_related('category').filter(manager=request.user, is_active=True).first()


def _categories(store):
    if store.category_id and store.category.sells_everything:
        qs = ProductCategory.objects.filter(Q(store__isnull=True) | Q(store=store))
    else:
        qs = ProductCategory.objects.filter(Q(store=store) | Q(store__isnull=True, store_category_id=store.category_id))
    return list(qs.order_by('name'))


def _find_category(categories, wanted):
    wanted = _plain(wanted)
    if not wanted:
        return None
    exact = [c for c in categories if _plain(c.name) == wanted]
    if exact:
        return exact[0]
    close = [c for c in categories if wanted in _plain(c.name) or _plain(c.name) in wanted]
    return close[0] if len(close) == 1 else None


def _number(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    text = re.sub(r'(?i)\s|fcfa|cfa|xaf|kg|f$', '', str(value).replace(' ', ''))
    text = text.replace(',', '.')
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        raise ValueError(str(value))


def _amount(value):
    """Montant en FCFA (pas de centimes) : « 4.500 » ou « 4,500 » se lisent 4 500, pas 4,5."""
    if isinstance(value, str) and re.fullmatch(r'\s*\d{1,3}([.,]\d{3})+\s*', value):
        value = re.sub(r'[.,]', '', value)
    return _number(value)


def _read_rows(upload):
    """Lignes du fichier en dictionnaires {colonne: valeur}, plus les colonnes non reconnues."""
    name = (upload.name or '').lower()
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValueError('Fichier trop lourd (2 Mo au plus). Coupez-le en plusieurs fichiers.')
    if name.endswith('.xlsx'):
        try:
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(raw), read_only=True, data_only=True).worksheets[0]
            table = [list(row) for row in sheet.iter_rows(values_only=True)]
        except Exception:
            raise ValueError("Ce fichier Excel ne peut pas être lu. Enregistrez-le en .xlsx ou en .csv puis réessayez.")
    elif name.endswith(('.csv', '.txt')):
        for encoding in ('utf-8-sig', 'cp1252'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("Les accents de ce fichier CSV ne peuvent pas être lus. Dans Excel, choisissez "
                             "« Enregistrer sous » puis « CSV UTF-8 », ou envoyez le fichier en .xlsx.")
        first = text.split('\n', 1)[0]
        delimiter = max((';', ',', '\t'), key=first.count)
        table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    else:
        raise ValueError('Format non pris en charge : envoyez un fichier Excel (.xlsx) ou CSV (.csv).')
    table = [row for row in table if any(str(cell or '').strip() for cell in row)]
    if not table:
        raise ValueError('Le fichier est vide.')
    header = [ALIASES.get(_plain(cell)) for cell in table[0]]
    unknown = [str(cell) for cell, key in zip(table[0], header) if cell not in (None, '') and key is None]
    missing = [k for k in REQUIRED if k not in header]
    if missing:
        labels = {'name': 'Nom', 'price': 'Prix', 'weight_kg': 'Poids (kg)', 'length_m': 'Longueur (m)'}
        raise ValueError('Colonnes obligatoires absentes : ' + ', '.join(labels[k] for k in missing)
                         + '. Partez du modèle à télécharger.')
    rows = []
    for cells in table[1:]:
        rows.append({key: cells[i] for i, key in enumerate(header) if key and i < len(cells)})
    if len(rows) > MAX_ROWS:
        raise ValueError(f'{len(rows)} lignes : {MAX_ROWS} au plus par fichier. Coupez-le en plusieurs fichiers.')
    return rows, unknown


def _flatten(errors):
    out = []
    for field, messages in (errors.items() if isinstance(errors, dict) else [('', errors)]):
        for message in (messages if isinstance(messages, list) else [messages]):
            out.append(str(message))
    return ' '.join(out)


def _prepare(row, store, categories, promotions):
    """(données pour le formulaire produit, alertes) ou ValueError avec la raison."""
    text = {k: str(v).strip() if v is not None else '' for k, v in row.items()}
    if not text.get('name'):
        raise ValueError('Nom manquant.')
    data = {'name': text['name'][:200], 'description': text.get('description', '')[:2000]}
    notes = []
    for key, label in (('price', 'Prix'), ('weight_kg', 'Poids'), ('length_m', 'Longueur'),
                       ('compare_price', 'Prix avant promo'), ('stock', 'Stock')):
        try:
            value = (_amount if key in ('price', 'compare_price') else _number)(row.get(key))
        except ValueError:
            raise ValueError(f'{label} illisible : « {row.get(key)} ».')
        if value is not None:
            data[key] = value
    if 'price' not in data:
        raise ValueError('Prix manquant.')
    if data['price'] <= 0:
        raise ValueError(f"Prix invalide : « {row.get('price')} ». Il doit être plus grand que 0.")
    data['stock'] = int(data.get('stock', 1))
    if data.get('compare_price') is not None and not promotions:
        data.pop('compare_price')
        notes.append('prix avant promo ignoré (promotions non incluses dans votre forfait)')
    if data.get('compare_price') is not None and data['compare_price'] <= data['price']:
        raise ValueError('Le prix avant promo doit être plus grand que le prix de vente.')
    if text.get('category'):
        category = _find_category(categories, text['category'])
        if category is None:
            notes.append(f"catégorie « {text['category']} » inconnue : produit créé sans catégorie")
        else:
            data['category'] = category.id
    attributes = {k: text[k][:100] for k in ('brand', 'model', 'color', 'sizes') if text.get(k)}
    if attributes:
        data['attributes'] = attributes
    for key in ('sku', 'barcode'):
        if text.get(key):
            data[key] = text[key][:100]
    return data, notes


@api_view(['GET'])
@permission_classes([permissions.AllowAny])
def import_template(request):
    """GET /api/v1/store/products/import/template/ : modèle CSV à remplir (s'ouvre dans Excel).

    Public (aucune donnée de la boutique) : l'application Android l'ouvre dans le navigateur du téléphone."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';')
    writer.writerow(TEMPLATE_HEADER)
    writer.writerows(TEMPLATE_ROWS)
    response = HttpResponse('﻿' + buffer.getvalue(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="modele-import-produits-gaboshop.csv"'
    return response


@api_view(['POST'])
@permission_classes([permissions.IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def import_products(request):
    """POST /api/v1/store/products/import/ (file, confirm=0|1) : aperçu, puis création des lignes valides."""
    store = _store(request)
    if store is None:
        return _error('no_store', "Aucune boutique active n'est liée à votre compte.",
                      'Créez ou réactivez votre boutique, puis réessayez.', status.HTTP_403_FORBIDDEN)
    try:
        plan = require(store, 'product_import')
    except PlanFeatureRequired as exc:
        return _error('plan_required', str(exc), 'Ouvrez « Mon forfait » et passez au forfait Pro.',
                      status.HTTP_403_FORBIDDEN)
    upload = request.FILES.get('file')
    if upload is None:
        return _error('no_file', 'Aucun fichier reçu.', 'Choisissez votre fichier Excel ou CSV, puis réessayez.')
    try:
        rows, unknown = _read_rows(upload)
    except ValueError as exc:
        return _error('bad_file', str(exc), 'Téléchargez le modèle, remplissez-le, puis réessayez.')

    confirm = str(request.data.get('confirm', '0')).lower() in ('1', 'true', 'oui')
    categories = _categories(store)
    promotions = has(plan, 'promotions')
    existing = {_plain(n) for n in store.products.values_list('name', flat=True)}
    max_products = getattr(plan, 'max_b2b_products', None) if store.is_b2b else plan.max_products
    room = None if max_products is None else max(0, max_products - store.products.count())

    results, created, seen = [], 0, set()
    for number, row in enumerate(rows, start=2):
        name = str(row.get('name') or '').strip()
        entry = {'row': number, 'name': name, 'status': 'ok', 'message': '', 'notes': []}
        results.append(entry)
        try:
            data, entry['notes'] = _prepare(row, store, categories, promotions)
        except ValueError as exc:
            entry.update(status='error', message=str(exc))
            continue
        # Ce qui sera mis en vente, pour que le commerçant vérifie les prix avant de confirmer.
        entry.update(price=int(data['price']), stock=data['stock'],
                     compare_price=int(data['compare_price']) if data.get('compare_price') else None)
        key = _plain(data['name'])
        if key in existing or key in seen:
            entry.update(status='skipped', message='Déjà dans votre boutique sous ce nom : non recréé.')
            continue
        serializer = ProductCreateSerializer(data=data, context={'request': request, 'store': store})
        if not serializer.is_valid():
            entry.update(status='error', message=_flatten(serializer.errors))
            continue
        if room is not None and len(seen) >= room:
            entry.update(status='error', message=SubscriptionChecker.product_limit_message(plan, max_products))
            continue
        try:
            SubscriptionChecker.check_can_add_food_product(store, serializer.validated_data.get('category'))
        except PermissionDenied as exc:
            entry.update(status='error', message=str(exc))
            continue
        seen.add(key)
        if confirm:
            with transaction.atomic():
                product = serializer.save(store=store)
            entry.update(status='created', product_id=product.id)
            created += 1

    counts = {s: sum(1 for r in results if r['status'] == s) for s in ('ok', 'created', 'skipped', 'error')}
    if confirm:
        message = (f"{created} produit{'s' if created > 1 else ''} créé{'s' if created > 1 else ''}."
                   + (f" {counts['error']} ligne(s) en erreur, non créée(s)." if counts['error'] else '')
                   + " Ajoutez maintenant leurs photos dans « Mes produits ».")
    else:
        message = (f"{counts['ok']} produit(s) prêt(s) à créer, {counts['error']} ligne(s) en erreur, "
                   f"{counts['skipped']} déjà présent(s). Rien n'est encore créé.")
    return Response({'success': True, 'data': {
        'confirmed': confirm, 'message': message, 'counts': counts, 'rows': results,
        'unknown_columns': unknown, 'categories': [c.name for c in categories],
    }})
